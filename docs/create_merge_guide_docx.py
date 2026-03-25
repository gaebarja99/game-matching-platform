# -*- coding: utf-8 -*-
"""Generate GameMatcher folder merge guide as Word .docx"""
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def add_para(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_after = Pt(6)
    return p


def main():
    # ASCII file name avoids Windows console / tool encoding issues; title is inside the document.
    out = Path(__file__).resolve().parent / "GameMatcher_folder_merge_guide.docx"
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"

    add_heading(doc, "GameMatcher 폴더 병합 가이드", 0)
    add_para(
        doc,
        "본 문서는 「GameMatcher (2)」 폴더와 「GameMatcher1_base」 폴더를 하나의 프로젝트로 안전하게 합치는 절차를 설명합니다. "
        "인프라와 설정이 더 많은 쪽인 GameMatcher1_base를 기준으로 두고, 다른 쪽의 고유 자산을 가져오는 방식을 권장합니다.",
    )

    add_heading(doc, "1. 사전 준비", 1)
    add_para(doc, "두 폴더 전체를 백업합니다(다른 드라이브나 클라우드에 복사본 보관).")
    add_para(doc, "Git을 쓰는 경우: GameMatcher1_base에서 병합 전에 브랜치를 만들거나 커밋합니다.")
    add_para(doc, "작업 중 Spring Boot·프론트 개발 서버가 실행 중이면 모두 종료합니다.")

    add_heading(doc, "2. 기준 폴더 정하기", 1)
    add_para(
        doc,
        "권장: GameMatcher1_base를 최종 프로젝트 루트로 둡니다(rtmp-server, secrets, config-examples, ffmpeg, frontend 등이 이미 포함되어 있음). "
        "GameMatcher (2)는 가져올 소스로만 사용합니다.",
    )

    add_heading(doc, "3. 문서·기획 자산(GameMatcher (2) 전용)", 1)
    add_para(
        doc,
        "Word 문서(ai 전직 분석, Phase4 LLM 가이드, 빌로란트 승리기여도 계획, 발로란트 API 목록, 데이터 가공 엔진 등)는 "
        "GameMatcher1_base\\docs 아래에 legacy 또는 planning 같은 하위 폴더를 만들어 복사합니다.",
    )
    add_para(doc, "기존 docs와 제목이 겹치면 파일명에 날짜나 출처를 붙여 구분합니다.")

    add_heading(doc, "4. 프론트엔드 폴더(gamematcher-frontend vs frontend)", 1)
    add_para(doc, "이름과 코드가 다를 수 있으므로 한쪽으로 무조건 덮어쓰지 않습니다.")
    add_para(
        doc,
        "권장 절차: (1) GameMatcher1_base의 frontend를 frontend_base 등으로 복사해 보관합니다. "
        "(2) GameMatcher (2)의 gamematcher-frontend를 별도 폴더로 복사합니다. "
        "(3) WinMerge, VS Code Compare Folders, 또는 git diff로 차이를 비교합니다. "
        "(4) 필요한 페이지·컴포넌트만 한쪽에서 다른 쪽으로 옮깁니다.",
    )
    add_para(
        doc,
        "package.json, vite.config 등 빌드 설정은 하나의 프로젝트로 통일하고, 의존성 버전을 맞춘 뒤 npm install을 다시 실행합니다.",
    )

    add_heading(doc, "5. Java 백엔드(src, pom.xml)", 1)
    add_para(
        doc,
        "pom.xml: 의존성·플러그인을 비교합니다. 한 파일을 기준으로 두고 다른 쪽의 <dependency>만 추가하는 방식이 안전합니다.",
    )
    add_para(
        doc,
        "src/main/java: 패키지가 같으면 컨트롤러·서비스·엔티티별로 비교합니다. 동일 클래스가 있으면 diff로 열어 메서드 단위로 병합합니다.",
    )
    add_para(doc, "src/main/resources: application*.properties, mapper XML, 정적 리소스도 동일하게 비교·병합합니다.")
    add_para(doc, "병합 후: mvnw.cmd clean compile(또는 mvn clean compile)으로 컴파일 오류를 확인합니다.")

    add_heading(doc, "6. 공통 폴더(docs, uploads, target, .idea)", 1)
    add_para(doc, "docs: 위 3절처럼 하위 폴더로 정리해 병합합니다.")
    add_para(doc, "uploads: 사용자 데이터가 있으면 백업 후 필요한 파일만 선택적으로 합칩니다.")
    add_para(doc, "target: 빌드 산출물이므로 병합하지 않고 삭제한 뒤 Maven으로 다시 빌드합니다.")
    add_para(doc, ".idea: 한쪽 설정만 쓰거나, 프로젝트를 다시 열어 IDE가 재생성하게 하면 충돌이 줄어듭니다.")

    add_heading(doc, "7. GameMatcher1_base에만 있는 항목", 1)
    add_para(
        doc,
        "rtmp-server, secrets, config-examples, donationSound, ffmpeg 등은 그대로 두고, "
        "GameMatcher (2)에 대응 파일이 없으면 변경하지 않습니다.",
    )

    add_heading(doc, "8. 병합 후 검증", 1)
    add_para(doc, "백엔드: README에 있는 프로필(예: mysql, oauth)로 애플리케이션을 기동해 봅니다.")
    add_para(doc, "프론트엔드: npm run dev 또는 프로젝트 스크립트로 UI를 확인합니다.")
    add_para(doc, "로그인, DB 연결, 주요 API를 스모크 테스트합니다.")

    add_heading(doc, "9. 충돌이 많을 때", 1)
    add_para(doc, "WinMerge, Meld, VS Code, Git merge 도구로 파일 단위로 해결합니다.")
    add_para(
        doc,
        "큰 기능 단위로 나누어 한쪽 버전에 태그를 찍은 뒤, 기능별로 가져오면(cherry-pick처럼) 위험이 줄어듭니다.",
    )

    doc.save(out)
    print(f"Created: {out}")


if __name__ == "__main__":
    main()
